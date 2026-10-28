import argparse
import csv
import re
from pathlib import Path

from openpyxl import load_workbook

FIELDS = [
    "name_zh", "name_en", "category_l1", "category_l2", "food_type", "dish_style",
    "region", "cuisine", "state", "edible_ratio", "serving_g", "common_portion_g",
    "energy_kcal", "protein_g", "fat_g", "carb_g", "fiber_g", "sodium_mg", "sugar_g",
    "ingredients", "taste_tags", "dietary_tags", "suitable_meals",
    "source_id", "source_url", "verified", "active",
]

COMPOSITE_WORDS = (
    "dish", "meal", "recipe", "pie", "pizza", "burger", "sandwich", "soup",
    "stew", "casserole", "curry", "ready", "sausage", "lasagne", "quiche",
)


def number(value):
    if value is None:
        return None
    text = str(value).strip().replace(",", "")
    if text.lower() in {"tr", "trace"}:
        return 0.0
    if text.upper() in {"N", "NA", "--", "-"}:
        return None
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    return float(match.group()) if match else None


def clean(value):
    return (value or "").strip()


def read_rows(workbook, sheet_name: str):
    sheet = workbook[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [clean(cell) for cell in rows[0]]
    data = []
    for row in rows[3:]:
        if all(cell is None or clean(cell) == "" for cell in row):
            continue
        data.append(row)
    return headers, data


def is_composite(name: str, description: str) -> bool:
    text = f"{name} {description}".lower()
    return any(word in text for word in COMPOSITE_WORDS)


def state_for(name: str, description: str, composite: bool):
    text = f"{name} {description}".lower()
    if "raw" in text:
        return "raw"
    if any(word in text for word in ("boiled", "cooked", "baked", "roasted", "fried", "grilled")):
        return "cooked"
    return "processed" if composite else "cooked"


def build(xlsx_path: Path, out_path: Path) -> None:
    workbook = load_workbook(xlsx_path, read_only=True, data_only=True)
    factors_headers, factors_rows = read_rows(workbook, "1.2 Factors")
    inorganics_headers, inorganics_rows = read_rows(workbook, "1.4 Inorganics")
    proximates_headers, proximates_rows = read_rows(workbook, "1.3 Proximates")

    edible_map = {}
    for row in factors_rows:
        code = clean(row[0])
        if code:
            edible_map[code] = number(row[7])

    sodium_map = {}
    for row in inorganics_rows:
        code = clean(row[0])
        if code:
            sodium_map[code] = number(row[7])

    index = {header: i for i, header in enumerate(proximates_headers)}
    rows = []
    for row in proximates_rows:
        code = clean(row[0])
        name = clean(row[1])
        if not code or not name:
            continue
        description = clean(row[2])
        composite = is_composite(name, description)
        edible = edible_map.get(code)
        if edible is None:
            edible = 1.0
        portion = 250 if composite else 100
        energy = number(row[index["Energy (kcal) (kcal)"]])
        if energy is None:
            continue
        values = {
            "name_zh": name,
            "name_en": name,
            "category_l1": "家常菜" if composite else "基础食材",
            "category_l2": clean(row[3]),
            "food_type": "dish" if composite else "ingredient",
            "dish_style": "家常菜" if composite else "基础食材",
            "region": "英国",
            "cuisine": "西式",
            "state": state_for(name, description, composite),
            "edible_ratio": edible,
            "serving_g": portion,
            "common_portion_g": portion,
            "energy_kcal": energy,
            "protein_g": number(row[index["Protein (g)"]]),
            "fat_g": number(row[index["Fat (g)"]]),
            "carb_g": number(row[index["Carbohydrate (g)"]]),
            "fiber_g": number(row[index["AOAC fibre (g)"]]),
            "sodium_mg": sodium_map.get(code),
            "sugar_g": number(row[index["Total sugars (g)"]]),
            "ingredients": "",
            "taste_tags": "",
            "dietary_tags": "",
            "suitable_meals": "",
            "source_id": code,
            "source_url": "https://www.gov.uk/government/publications/composition-of-foods-integrated-dataset-cofid",
            "verified": "true",
            "active": "false",
        }
        rows.append(values)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(rows)
    print(f"wrote {len(rows)} foods to {out_path}")


def main() -> None:
    parser = argparse.ArgumentParser(description="把 CoFID Excel 转换成规范营养表")
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()
    build(args.xlsx, args.out)


if __name__ == "__main__":
    main()
