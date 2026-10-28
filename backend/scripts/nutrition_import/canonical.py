import json
import re
from pathlib import Path

from app.db.session import SessionLocal
from app.models.food import Food, ImportBatch
from app.services.food_text import infer_food_type, infer_region, normalize_name
from app.services.llm import foods_data

NUTRIENT_FIELDS = (
    "energy_kcal",
    "protein_g",
    "fat_g",
    "carb_g",
    "fiber_g",
    "sodium_mg",
    "sugar_g",
)
LIST_FIELDS = ("ingredients", "taste_tags", "dietary_tags", "suitable_meals")
BOOL_FIELDS = ("verified", "active")
CATEGORY_STYLE = {
    "主食": "主食",
    "荤菜": "家常菜",
    "素菜": "家常菜",
    "汤粥": "汤粥",
    "早餐": "早餐",
    "小吃": "小吃",
    "水果": "单品",
    "饮品": "饮品",
}

FIELD_ALIASES = {
    "name": "name_zh",
    "name_zh": "name_zh",
    "名称": "name_zh",
    "菜名": "name_zh",
    "样品名称": "name_zh",
    "食品名称": "name_zh",
    "name_en": "name_en",
    "英文名": "name_en",
    "category": "category_l1",
    "category_l1": "category_l1",
    "类别": "category_l1",
    "食品分类": "category_l1",
    "category_l2": "category_l2",
    "子类别": "category_l2",
    "food_type": "food_type",
    "类型": "food_type",
    "dish_style": "dish_style",
    "菜品类型": "dish_style",
    "region": "region",
    "地区": "region",
    "cuisine": "cuisine",
    "菜系": "cuisine",
    "state": "state",
    "状态": "state",
    "edible_ratio": "edible_ratio",
    "可食部": "edible_ratio",
    "waste_ratio": "waste_ratio",
    "废弃率": "waste_ratio",
    "serving_g": "serving_g",
    "每份克重": "serving_g",
    "份量": "serving_g",
    "common_portion_g": "common_portion_g",
    "标准份": "common_portion_g",
    "常见份量": "common_portion_g",
    "energy_kcal": "energy_kcal",
    "kcal_per_100g": "energy_kcal",
    "calories": "energy_kcal",
    "热量": "energy_kcal",
    "能量": "energy_kcal",
    "每100克热量": "energy_kcal",
    "protein_g": "protein_g",
    "protein": "protein_g",
    "蛋白质": "protein_g",
    "fat_g": "fat_g",
    "fat": "fat_g",
    "脂肪": "fat_g",
    "carb_g": "carb_g",
    "carb": "carb_g",
    "carbohydrate": "carb_g",
    "碳水": "carb_g",
    "碳水化合物": "carb_g",
    "fiber_g": "fiber_g",
    "fiber": "fiber_g",
    "纤维": "fiber_g",
    "sodium_mg": "sodium_mg",
    "sodium": "sodium_mg",
    "钠": "sodium_mg",
    "sugar_g": "sugar_g",
    "sugar": "sugar_g",
    "糖": "sugar_g",
    "ingredients": "ingredients",
    "食材": "ingredients",
    "原料": "ingredients",
    "taste_tags": "taste_tags",
    "口味": "taste_tags",
    "dietary_tags": "dietary_tags",
    "饮食标签": "dietary_tags",
    "suitable_meals": "suitable_meals",
    "适用餐次": "suitable_meals",
    "source_id": "source_id",
    "来源编号": "source_id",
    "license": "license",
    "许可": "license",
    "source_url": "source_url",
    "来源链接": "source_url",
    "version": "version",
    "版本": "version",
    "verified": "verified",
    "已校验": "verified",
    "active": "active",
    "启用": "active",
}


def parse_number(value) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    text = re.sub(r"[^0-9.\-]", "", text)
    if text in ("", "-", ".", "-."):
        return None
    try:
        return float(text)
    except ValueError:
        return None


def parse_bool(value) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    return str(value).strip().lower() in {"1", "true", "yes", "y", "是", "已校验", "启用"}


def parse_list(value) -> list:
    if value is None or value == "":
        return []
    if isinstance(value, list):
        return value
    text = str(value).strip()
    if text.startswith("["):
        try:
            parsed = json.loads(text)
            return parsed if isinstance(parsed, list) else []
        except json.JSONDecodeError:
            pass
    return [part.strip() for part in re.split(r"[|,，、;；]", text) if part.strip()]


def normalize_row(raw: dict, source: str, version: str, license: str, source_type: str) -> dict:
    row = {}
    for key, value in raw.items():
        if key is None:
            continue
        target = FIELD_ALIASES.get(str(key).strip(), str(key).strip())
        if target not in row or row[target] in (None, ""):
            row[target] = value

    name = str(row.get("name_zh") or "").strip()
    if not name:
        raise ValueError("missing name_zh")
    category = str(row.get("category_l1") or "其他").strip()
    cuisine = str(row.get("cuisine") or "家常菜").strip()
    food_type = str(row.get("food_type") or infer_food_type(category)).strip()
    state = str(row.get("state") or "cooked").strip()
    waste = parse_number(row.get("waste_ratio"))
    edible = parse_number(row.get("edible_ratio"))
    if edible is None and waste is not None:
        edible = 1 - waste / 100 if waste > 1 else 1 - waste
    if edible is None:
        edible = 1.0
    serving = parse_number(row.get("serving_g"))
    portion = parse_number(row.get("common_portion_g")) or serving or 100
    nutrients = {field: parse_number(row.get(field)) for field in NUTRIENT_FIELDS}
    return {
        "source": source,
        "source_id": str(row.get("source_id") or name).strip(),
        "name_zh": name,
        "name_en": (str(row.get("name_en")).strip() if row.get("name_en") else None),
        "name_norm": normalize_name(name),
        "food_type": food_type,
        "dish_style": str(row.get("dish_style") or CATEGORY_STYLE.get(category, "家常菜")).strip(),
        "region": str(row.get("region") or infer_region(cuisine)).strip(),
        "cuisine": cuisine,
        "category_l1": category,
        "category_l2": (str(row.get("category_l2")).strip() if row.get("category_l2") else None),
        "state": state,
        "edible_ratio": edible,
        "serving_g": serving or portion,
        "common_portion_g": portion,
        "energy_kcal": nutrients["energy_kcal"],
        "protein_g": nutrients["protein_g"],
        "fat_g": nutrients["fat_g"],
        "carb_g": nutrients["carb_g"],
        "fiber_g": nutrients["fiber_g"],
        "sodium_mg": nutrients["sodium_mg"],
        "sugar_g": nutrients["sugar_g"],
        "ingredients": parse_list(row.get("ingredients")),
        "suitable_meals": parse_list(row.get("suitable_meals")),
        "taste_tags": parse_list(row.get("taste_tags")),
        "dietary_tags": parse_list(row.get("dietary_tags")),
        "source_type": source_type,
        "license": license,
        "source_url": (str(row.get("source_url")).strip() if row.get("source_url") else None),
        "version": str(row.get("version") or version),
        "verified": parse_bool(row.get("verified")),
        "active": parse_bool(row.get("active")) if row.get("active") is not None else True,
    }


def validate_row(values: dict) -> list[str]:
    errors = []
    energy = values["energy_kcal"]
    protein = values["protein_g"] or 0
    fat = values["fat_g"] or 0
    carb = values["carb_g"] or 0
    if energy is None or energy < 0 or energy > 900:
        errors.append("energy_kcal out of range")
    if not (0 <= protein <= 100 and 0 <= fat <= 105 and 0 <= carb <= 105):
        errors.append("macros out of range")
    if not (0 < values["common_portion_g"] <= 3000):
        errors.append("common_portion_g out of range")
    if not (0 <= values["edible_ratio"] <= 1):
        errors.append("edible_ratio out of range")
    if values["food_type"] not in {"ingredient", "single_food", "dish", "light_meal", "packaged", "beverage"}:
        errors.append("invalid food_type")
    if values["state"] not in {"raw", "cooked", "processed", "ready_to_eat"}:
        errors.append("invalid state")
    return errors


def validate_warnings(values: dict) -> list[str]:
    warnings = []
    energy = values["energy_kcal"]
    protein = values["protein_g"] or 0
    fat = values["fat_g"] or 0
    carb = values["carb_g"] or 0
    estimated = 4 * protein + 9 * fat + 4 * carb
    if energy and abs(estimated - energy) > max(60, energy * 0.35):
        warnings.append(f"macro-energy mismatch ({energy} vs {estimated:.0f})")
    missing = [field for field in NUTRIENT_FIELDS if values.get(field) is None]
    if missing:
        warnings.append(f"missing nutrients: {','.join(missing)}")
    return warnings


def import_rows(
    rows: list[dict],
    source: str,
    version: str,
    license: str,
    source_type: str = "official",
    dry_run: bool = True,
    force: bool = False,
    checksum: str | None = None,
) -> dict:
    normalized = []
    errors = []
    warnings = []
    for index, raw in enumerate(rows, start=2):
        try:
            values = normalize_row(raw, source, version, license, source_type)
            row_errors = validate_row(values)
            row_warnings = validate_warnings(values)
        except Exception as exc:
            errors.append({"row": index, "error": str(exc)})
            continue
        if row_warnings:
            warnings.append({"row": index, "name": values.get("name_zh"), "warnings": row_warnings})
        if row_errors:
            errors.append({"row": index, "name": values.get("name_zh"), "errors": row_errors})
            continue
        normalized.append(values)

    unique = {}
    for values in normalized:
        key = (values["source"], values["source_id"])
        if key in unique:
            warnings.append({
                "row": None,
                "name": values["name_zh"],
                "warnings": [f"duplicate source_id {values['source_id']} ignored"],
            })
            continue
        unique[key] = values
    normalized = list(unique.values())

    if dry_run:
        return {
            "source": source,
            "total": len(rows),
            "valid": len(normalized),
            "errors": errors,
            "warnings": warnings,
            "dry_run": True,
        }

    if errors:
        return {
            "source": source,
            "total": len(rows),
            "valid": len(normalized),
            "errors": errors,
            "warnings": warnings,
            "imported": 0,
            "dry_run": False,
        }

    db = SessionLocal()
    inserted = updated = 0
    try:
        for values in normalized:
            food = (
                db.query(Food)
                .filter(Food.source == source, Food.source_id == values["source_id"])
                .first()
            )
            if food is None:
                db.add(Food(**values))
                inserted += 1
            elif force:
                for key, value in values.items():
                    setattr(food, key, value)
                updated += 1
        db.add(ImportBatch(
            source=source,
            version=version,
            license=license,
            row_count=len(normalized),
            checksum=checksum,
        ))
        db.commit()
    except Exception:
        db.rollback()
        raise
    finally:
        db.close()

    foods_data.reload()
    return {
        "source": source,
        "total": len(rows),
        "valid": len(normalized),
        "imported": inserted,
        "updated": updated,
        "errors": errors,
        "warnings": warnings,
        "dry_run": False,
    }


def load_csv(path: Path) -> list[dict]:
    import csv

    with open(path, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def load_excel(path: Path) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("读取 Excel 需要先安装 openpyxl：pip install openpyxl") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(cell).strip() if cell is not None else "" for cell in rows[0]]
    result = []
    for row in rows[1:]:
        if all(cell is None or str(cell).strip() == "" for cell in row):
            continue
        result.append({headers[i]: row[i] for i in range(min(len(headers), len(row)))})
    workbook.close()
    return result
